"""
extraction.py
-------------
Job: take an uploaded file (PDF or DOCX) and pull out plain text + any
tables it finds. This is the "multi-format data extraction" part of
the resume bullet.

Why separate file? Keeps parsing logic isolated so you can add more
formats (images with OCR, CSV, etc.) later without touching anything
else.
"""

from pypdf import PdfReader
from docx import Document
import io


def extract_from_pdf(file_bytes: bytes) -> dict:
    """Returns {'text': str, 'tables': list, 'num_pages': int}"""
    reader = PdfReader(io.BytesIO(file_bytes))
    full_text = []
    for page in reader.pages:
        page_text = page.extract_text() or ""
        full_text.append(page_text)

    # NOTE: pypdf doesn't do great table extraction out of the box.
    # For a resume project, plain text extraction is enough to start.
    # If you want proper table detection later, look into `camelot-py`
    # or `pdfplumber` as a follow-up improvement.
    return {
        "text": "\n".join(full_text),
        "tables": [],
        "num_pages": len(reader.pages),
    }


def extract_from_docx(file_bytes: bytes) -> dict:
    """Returns {'text': str, 'tables': list[list[list[str]]]}"""
    doc = Document(io.BytesIO(file_bytes))

    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

    tables = []
    for table in doc.tables:
        rows = []
        for row in table.rows:
            rows.append([cell.text for cell in row.cells])
        tables.append(rows)

    return {
        "text": "\n".join(paragraphs),
        "tables": tables,
        "num_pages": None,
    }


def extract_document(filename: str, file_bytes: bytes) -> dict:
    """Router: picks the right extractor based on file extension."""
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return extract_from_pdf(file_bytes)
    elif lower.endswith(".docx"):
        return extract_from_docx(file_bytes)
    else:
        raise ValueError(f"Unsupported file type: {filename}")
